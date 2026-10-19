"""File conversions: anything -> .docx -> .pdf.

Word (via COM automation) gives the best fidelity for documents with images and Arabic text.
LibreOffice is used as a fallback when Word is not available.
"""
from __future__ import annotations

import logging
import re
import shutil
import subprocess
import threading
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

log = logging.getLogger(__name__)
_word_lock = threading.Lock()

WORD_EXTS = {".docx", ".doc", ".rtf", ".odt"}
TEXT_EXTS = {".md", ".markdown", ".txt"}
SUPPORTED = WORD_EXTS | TEXT_EXTS | {".pdf"}

WD_EXPORT_PDF = 17
WD_EXPORT_HEADING_BOOKMARKS = 1
WD_FORMAT_DOCX = 16


def _soffice() -> str | None:
    for p in (shutil.which("soffice"), r"C:\Program Files\LibreOffice\program\soffice.exe"):
        if p and Path(p).exists():
            return p
    return None


def _with_word(fn):
    """Run fn(word_app) inside a private, invisible Word instance."""
    import pythoncom
    import win32com.client

    with _word_lock:
        pythoncom.CoInitialize()
        word = None
        try:
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            word.DisplayAlerts = 0
            return fn(word)
        finally:
            if word is not None:
                try:
                    word.Quit(SaveChanges=0)
                except Exception:
                    pass
            pythoncom.CoUninitialize()


def word_to_pdf(src: Path, out_pdf: Path) -> Path:
    out_pdf.parent.mkdir(parents=True, exist_ok=True)
    try:
        def run(word):
            d = word.Documents.Open(str(src.resolve()), ReadOnly=True, AddToRecentFiles=False,
                                    ConfirmConversions=False, Visible=False)
            try:
                d.ExportAsFixedFormat(OutputFileName=str(out_pdf.resolve()), ExportFormat=WD_EXPORT_PDF,
                                      OpenAfterExport=False, OptimizeFor=0,
                                      CreateBookmarks=WD_EXPORT_HEADING_BOOKMARKS, DocStructureTags=True)
            finally:
                d.Close(SaveChanges=0)
        _with_word(run)
        return out_pdf
    except ImportError:
        pass
    except Exception as e:
        log.warning("Word conversion failed for %s: %s", src, e)
        if not _soffice():
            raise
    soffice = _soffice()
    if not soffice:
        raise RuntimeError("Neither Microsoft Word nor LibreOffice is available for PDF conversion")
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_pdf.parent), str(src)],
                   check=True, timeout=300, capture_output=True)
    produced = out_pdf.parent / (src.stem + ".pdf")
    if produced != out_pdf:
        produced.replace(out_pdf)
    return out_pdf


def to_docx(src: Path, out_docx: Path) -> Path:
    """Convert anything Word can open (.doc, .rtf, .odt, text PDFs via PDF reflow) into .docx."""
    def run(word):
        d = word.Documents.Open(str(src.resolve()), ConfirmConversions=False, ReadOnly=True,
                                AddToRecentFiles=False, Visible=False)
        try:
            d.SaveAs2(str(out_docx.resolve()), FileFormat=WD_FORMAT_DOCX)
        finally:
            d.Close(SaveChanges=0)
    _with_word(run)
    return out_docx


HTML_TEMPLATE = """<!doctype html><html><head><meta charset="utf-8"><title>{title}</title>
<style>
body {{ font-family: Calibri, Arial, sans-serif; font-size: 11pt; line-height: 1.35; }}
h1 {{ font-size: 22pt; color: #0B3F9E; }} h2 {{ font-size: 16pt; color: #0B3F9E; }} h3 {{ font-size: 13pt; }}
code, pre {{ font-family: Consolas, monospace; font-size: 9.5pt; }} pre {{ background: #F3F5F9; padding: 6pt; }}
blockquote {{ border-left: 3pt solid #9DB6E8; margin-left: 0; padding-left: 9pt; color: #444; }}
mark {{ background: #FFF3A3; }} table {{ border-collapse: collapse; }} td, th {{ border: 1px solid #999; padding: 3pt 6pt; }}
</style></head><body><h1>{title}</h1>{body}</body></html>"""


def html_export(html: str, title: str, out: Path, fmt: str) -> Path:
    """Render editor HTML to .docx or .pdf with Word (LibreOffice fallback), keeping formatting and font sizes."""
    import html as htmllib

    out.parent.mkdir(parents=True, exist_ok=True)
    src = out.with_suffix(".src.html")
    # Word ignores <mark>; its own highlight property survives the conversion
    html = re.sub(r"<mark\b[^>]*>", '<span style="background:yellow;mso-highlight:yellow">', html)
    html = html.replace("</mark>", "</span>")
    src.write_text(HTML_TEMPLATE.format(title=htmllib.escape(title), body=html), encoding="utf-8")
    try:
        def run(word):
            d = word.Documents.Open(str(src.resolve()), ConfirmConversions=False, ReadOnly=True,
                                    AddToRecentFiles=False, Visible=False)
            try:
                if fmt == "pdf":
                    d.ExportAsFixedFormat(OutputFileName=str(out.resolve()), ExportFormat=WD_EXPORT_PDF,
                                          OpenAfterExport=False, OptimizeFor=0,
                                          CreateBookmarks=WD_EXPORT_HEADING_BOOKMARKS)
                else:
                    d.SaveAs2(str(out.resolve()), FileFormat=WD_FORMAT_DOCX)
            finally:
                d.Close(SaveChanges=0)
        _with_word(run)
    except Exception as e:
        log.warning("Word export failed (%s), trying LibreOffice", e)
        soffice = _soffice()
        if not soffice:
            raise RuntimeError("Neither Microsoft Word nor LibreOffice is available for export") from e
        target = "pdf" if fmt == "pdf" else 'docx:"MS Word 2007 XML"'
        subprocess.run([soffice, "--headless", "--convert-to", target, "--outdir", str(out.parent), str(src)],
                       check=True, timeout=300, capture_output=True)
        (out.parent / (src.stem + "." + fmt)).replace(out)
    finally:
        src.unlink(missing_ok=True)
    return out


def html_to_text(html: str) -> str:
    """Plain text from editor HTML: block elements become line breaks, list items get bullets."""
    import html as htmllib

    t = re.sub(r"(?is)<(script|style)\b.*?</\1>", "", html)
    t = re.sub(r"(?i)<li[^>]*>", "\n• ", t)
    t = re.sub(r"(?i)<br\s*/?>", "\n", t)
    t = re.sub(r"(?i)</(p|h[1-6]|li|blockquote|pre|tr|div)>", "\n", t)
    t = re.sub(r"(?i)</t[dh]>", "\t", t)
    t = htmllib.unescape(re.sub(r"<[^>]+>", "", t))
    return re.sub(r"\n{3,}", "\n\n", "\n".join(line.rstrip() for line in t.splitlines())).strip() + "\n"


# ---------------------------------------------------------------------------
# Markdown / plain text -> .docx
# ---------------------------------------------------------------------------

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|(?<![\w*])\*[^*\s][^*]*\*(?![\w*])|\[[^\]]+\]\([^)]+\))")
IMAGE_RE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)$|^\[\[IMG:([^\]]+)\]\]$")
NAVY = RGBColor(0x0B, 0x3F, 0x9E)


def _add_inline(par, text: str):
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        elif part.startswith("[") and "](" in part:
            par.add_run(part[1:part.index("](")]).underline = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def _style_book(d):
    """Typography shared by every generated chapter."""
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color in (("Title", 26, NAVY), ("Heading 1", 17, NAVY), ("Heading 2", 13.5, NAVY),
                              ("Heading 3", 12, RGBColor(0x33, 0x33, 0x33))):
        st = d.styles[name]
        st.font.name, st.font.size, st.font.color.rgb, st.font.bold = "Calibri", Pt(size), color, True
        st.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True


def _add_image(d, path: Path, caption: str):
    from docx.shared import Inches
    width = 5.6
    try:  # cap the height so tall screenshots do not take a whole page
        import pymupdf
        pix = pymupdf.Pixmap(str(path))
        if pix.width and pix.height:
            width = min(5.6, 3.4 * pix.width / pix.height, pix.width / 96)
    except Exception:
        pass
    try:
        d.add_picture(str(path), width=Inches(max(width, 1.5)))
    except Exception:
        try:
            d.add_picture(str(path))
        except Exception as e:
            log.warning("could not insert image %s: %s", path, e)
            return
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        cap = d.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(caption)
        r.italic, r.font.size, r.font.color.rgb = True, Pt(9), RGBColor(0x55, 0x55, 0x55)


def _split_row(line: str) -> list[str]:
    line = line.strip().strip("|")
    cells = re.split(r"(?<!\\)\|", line)
    return [c.strip().replace("\\|", "|") for c in cells]


BLOCK_START = re.compile(r"^\s*(#{1,6}\s|```|\||!\[|\[\[IMG:|<!--|>|[-*]\s|\d+\.\s|(-{3,}|\*{3,})\s*$)")


def _unwrap(lines: list[str]) -> list[str]:
    """Join hard-wrapped lines into logical paragraphs / list items (code blocks and tables are left alone)."""
    out: list[str] = []
    in_code = False
    for line in lines:
        s = line.strip()
        if s.startswith("```"):
            in_code = not in_code
            out.append(line)
            continue
        if in_code:
            out.append(line)
            continue
        prev = out[-1].strip() if out else ""
        joinable_prev = prev and not re.match(r"^(#{1,6}\s|```|\||!\[|\[\[IMG:|<!--|(-{3,}|\*{3,})\s*$)", prev)
        if s and not BLOCK_START.match(line) and joinable_prev:
            out[-1] = out[-1].rstrip() + " " + s
        else:
            out.append(line)
    return out


def markdown_to_docx(text: str, out_docx: Path, title: str | None = None, image_dir: Path | None = None) -> Path:
    d = docx.Document()
    _style_book(d)
    if title and not text.lstrip().startswith("# "):
        d.add_heading(title, level=0)

    lines = _unwrap(text.splitlines())
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            p = d.add_paragraph()
            run = p.add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            p.paragraph_format.left_indent = Pt(12)
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{3,}", lines[i + 1]):
            header = _split_row(stripped)
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_split_row(lines[i]))
                i += 1
            table = d.add_table(rows=1 + len(rows), cols=len(header))
            table.style = "Table Grid"
            for c, h in enumerate(header):
                cell = table.rows[0].cells[c]
                cell.text = ""
                cell.paragraphs[0].add_run(h.replace("**", "")).bold = True
            for r, row in enumerate(rows, start=1):
                for c in range(len(header)):
                    cell = table.rows[r].cells[c]
                    cell.text = ""
                    _add_inline(cell.paragraphs[0], row[c] if c < len(row) else "")
            d.add_paragraph()
            continue
        if stripped.startswith("<!--"):      # hidden markers, such as where a workspace note was inserted
            i += 1
            continue
        img = IMAGE_RE.match(stripped)
        if img and image_dir is not None:
            ref = (img.group(2) or img.group(3) or "").strip()
            for cand in (image_dir / ref, image_dir / "images" / ref):
                if cand.exists():
                    _add_image(d, cand, (img.group(1) or "").strip())
                    break
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", stripped)
        indent = len(line) - len(line.lstrip(" "))
        if m:
            level = len(m.group(1))
            d.add_heading(m.group(2).strip(), level=0 if level == 1 else level - 1)
        elif re.match(r"^[-*]\s+\[[ xX]\]\s+", stripped):
            done = stripped[3] in "xX"
            _add_inline(d.add_paragraph(style="List Bullet"),
                        ("☑ " if done else "☐ ") + re.sub(r"^[-*]\s+\[[ xX]\]\s+", "", stripped))
        elif re.match(r"^[-*]\s+", stripped):
            p = d.add_paragraph(style="List Bullet 2" if indent >= 2 else "List Bullet")
            p.paragraph_format.space_after = Pt(2)
            _add_inline(p, re.sub(r"^[-*]\s+", "", stripped))
        elif re.match(r"^\d+\.\s+", stripped):
            # literal numbers: Word's "List Number" style would keep counting across separate lists
            from docx.shared import Inches
            num, body = re.match(r"^(\d+)\.\s+(.*)", stripped).groups()
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5 if indent >= 2 else 0.3)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.space_after = Pt(2)
            p.add_run(f"{num}.\t")
            _add_inline(p, body)
        elif re.match(r"^(-{3,}|\*{3,})$", stripped):
            pass  # horizontal rule: section spacing already separates content
        elif stripped.startswith(">"):
            p = d.add_paragraph(style="Intense Quote")
            _add_inline(p, stripped.lstrip("> "))
        elif stripped:
            _add_inline(d.add_paragraph(), stripped)
        i += 1

    for p in d.paragraphs:
        if p.style.name == "Title":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    out_docx.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out_docx))
    return out_docx
